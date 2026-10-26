from docx import Document

from traka_automation.enrollments.enrollment.enrollment_form import (
    generate_docx_enrollment_form,
)
from traka_automation.enrollments.models import EnrollmentWebForm, Participant
from traka_automation.util.dutch_date import dutch_date

# Fields that are not rendered as plain text in the form.
NON_TEXT_FIELDS = {"camp", "photo"}


def _document_text(docx) -> str:
    document = Document(docx)
    parts = [p.text for p in document.paragraphs]
    stack = list(document.tables)
    while stack:
        table = stack.pop()
        for cell in (c for row in table.rows for c in row.cells):
            parts.extend(p.text for p in cell.paragraphs)
            stack.extend(cell.tables)
    return "\n".join(parts)


def test_all_participant_fields_in_enrollment_form(
    example_enrollment_web_form: EnrollmentWebForm, tmp_path
):
    participant = example_enrollment_web_form.participants[0]
    path = tmp_path / "form.docx"
    generate_docx_enrollment_form(participant).save(str(path))
    text = _document_text(path)

    fields = set(Participant.model_fields) - NON_TEXT_FIELDS
    assert fields, "Participant has no text fields to check"

    missing = []
    for field in sorted(fields):
        value = getattr(participant, field)
        assert str(value), f"fixture has empty {field}; test would be vacuous"
        expected = dutch_date(value) if field == "birth_date" else str(value)
        if expected not in text:
            missing.append(field)

    assert not missing, f"Participant fields missing from enrollment form: {missing}"
